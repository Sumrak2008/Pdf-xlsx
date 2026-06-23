"""End-to-end tests over the 16 synthetic-PDF scenarios in tools/synthetic_pdfs.py.

Each test generates its own scenario PDF into tmp_path and runs it through
the real pipeline (and, where relevant, all the way to the written XLSX),
so this suite is the single place that proves the whole stack - not just
an isolated module - behaves correctly for each required case from the
specification: text/scanned/mixed documents, Russian/English/mixed text,
bordered/borderless tables, rotated pages (both /Rotate-flag and
baked-in), low-quality scans, uncertain-cell highlighting, and the two
required error paths (corrupt/encrypted files).
"""

from __future__ import annotations

import openpyxl
import pytest

from pdfxlsx.core import paths
from pdfxlsx.core.config import ConversionSettings
from pdfxlsx.core.errors import CorruptPdfError, EncryptedPdfError
from pdfxlsx.core.pipeline import convert_document, resolve_output_paths, run_conversion
from synthetic_pdfs import (
    scenario_01_text_table_russian,
    scenario_02_text_table_english,
    scenario_03_mixed_language_table,
    scenario_04_borderless_table_direct,
    scenario_05_merged_header_table,
    scenario_06_multi_page_with_blank,
    scenario_07_table_with_paragraph_text,
    scenario_08_scanned_table_clean,
    scenario_09_scanned_table_borderless,
    scenario_10_scanned_table_for_uncertain_highlighting,
    scenario_11_low_quality_noisy_scan,
    scenario_12_rotated_text_page_rotate_flag,
    scenario_13_rotated_scanned_page_baked_in,
    scenario_14_mixed_text_and_scanned_pages,
    scenario_15_corrupt_pdf,
    scenario_16_encrypted_pdf,
)

requires_tesseract = pytest.mark.skipif(
    paths.tesseract_executable() is None, reason="tesseract not available"
)


def _cells_by_pos(table):
    return {(c.row, c.col): c.display_text for c in table.cells}


def test_scenario_01_text_table_russian(tmp_path):
    pdf_path = scenario_01_text_table_russian(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    page = result.pages[0]
    assert page.page_kind == "text"
    assert page.extraction_method == "direct"
    values = _cells_by_pos(result.tables[0])
    assert values[(0, 0)] == "Наименование"
    assert values[(1, 0)] == "Болт М6х20"
    assert values[(1, 1)] == "150"


def test_scenario_02_text_table_english(tmp_path):
    pdf_path = scenario_02_text_table_english(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    page = result.pages[0]
    assert page.extraction_method == "direct"
    values = _cells_by_pos(result.tables[0])
    assert values[(0, 0)] == "Item"
    assert values[(1, 0)] == "Bolt M6x20"


def test_scenario_03_mixed_language_table(tmp_path):
    pdf_path = scenario_03_mixed_language_table(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    values = _cells_by_pos(result.tables[0])
    assert values[(1, 0)] == "A-001"
    assert values[(1, 1)] == "Болт М6х20 ГОСТ 7798-70"
    assert values[(1, 2)] == "USD"


def test_scenario_04_borderless_table_direct(tmp_path):
    pdf_path = scenario_04_borderless_table_direct(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    page = result.pages[0]
    assert page.extraction_method == "direct"
    assert len(result.tables) == 1
    assert result.tables[0].source == "text"
    values = _cells_by_pos(result.tables[0])
    assert values[(0, 1)] == "Наименование"
    assert values[(1, 0)] == "001"


def test_scenario_05_merged_header_table(tmp_path):
    pdf_path = scenario_05_merged_header_table(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    table = result.tables[0]
    header = next(c for c in table.cells if c.row == 0)
    assert header.col_span == 3
    assert header.display_text == "Сведения о товаре"


def test_scenario_06_multi_page_with_blank(tmp_path):
    pdf_path = scenario_06_multi_page_with_blank(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    assert result.page_count == 3
    assert [p.page_kind for p in result.pages] == ["text", "blank", "text"]
    assert [p.extraction_method for p in result.pages] == ["direct", "blank", "direct"]
    assert len(result.tables) == 2


def test_scenario_07_table_with_paragraph_text(tmp_path):
    pdf_path = scenario_07_table_with_paragraph_text(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    page = result.pages[0]
    assert page.extraction_method == "direct"
    assert any("Сопроводительное письмо" in p for p in page.plain_text_paragraphs)
    values = _cells_by_pos(result.tables[0])
    assert values[(1, 0)] == "Товар А"


@requires_tesseract
def test_scenario_08_scanned_table_clean(tmp_path):
    pdf_path = scenario_08_scanned_table_clean(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    page = result.pages[0]
    assert page.page_kind == "scanned"
    assert page.extraction_method == "ocr"
    values = _cells_by_pos(result.tables[0])
    assert any("Товар А" in v for v in values.values())
    assert any("10" in v for v in values.values())


@requires_tesseract
def test_scenario_09_scanned_table_borderless(tmp_path):
    pdf_path = scenario_09_scanned_table_borderless(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    assert len(result.tables) == 1
    assert result.tables[0].source == "ocr-text"
    values = _cells_by_pos(result.tables[0])
    assert any("Болт" in v for v in values.values())


@requires_tesseract
def test_scenario_10_uncertain_cells_highlighted_in_xlsx(tmp_path):
    pdf_path = scenario_10_scanned_table_for_uncertain_highlighting(tmp_path)
    settings = ConversionSettings(ocr_confidence_threshold=99)
    result = run_conversion(str(pdf_path), settings)

    uncertain = [c for t in result.tables for c in t.cells if c.is_uncertain]
    assert uncertain
    for cell in uncertain:
        assert cell.comment is not None

    xlsx_path, _ = resolve_output_paths(str(pdf_path), settings)
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.worksheets[0]
    flagged_fill_found = any(
        cell.fill.start_color.rgb in ("FFFFFF00", "00FFFF00")
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert flagged_fill_found


@requires_tesseract
def test_scenario_11_low_quality_noisy_scan_does_not_crash(tmp_path):
    pdf_path = scenario_11_low_quality_noisy_scan(tmp_path)
    settings = ConversionSettings()
    result = run_conversion(str(pdf_path), settings)

    assert result.pages[0].extraction_method == "ocr"
    xlsx_path, _ = resolve_output_paths(str(pdf_path), settings)
    assert xlsx_path.exists()


@requires_tesseract
def test_scenario_12_rotated_text_page_rotate_flag(tmp_path):
    pdf_path = scenario_12_rotated_text_page_rotate_flag(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    page = result.pages[0]
    assert page.extraction_method == "ocr"
    assert page.rotation_degrees == 90
    values = _cells_by_pos(result.tables[0])
    assert any("Товар А" in v for v in values.values())
    assert any("150" in v for v in values.values())


@requires_tesseract
def test_scenario_13_rotated_scanned_page_baked_in(tmp_path):
    pdf_path = scenario_13_rotated_scanned_page_baked_in(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    page = result.pages[0]
    assert page.extraction_method == "ocr"
    assert page.rotation_degrees != 0
    values = _cells_by_pos(result.tables[0])
    assert any("Товар А" in v for v in values.values())
    assert any("Товар В" in v for v in values.values())


@requires_tesseract
def test_scenario_14_mixed_text_and_scanned_pages(tmp_path):
    pdf_path = scenario_14_mixed_text_and_scanned_pages(tmp_path)
    result = convert_document(str(pdf_path), ConversionSettings())

    assert result.page_count == 2
    assert [p.extraction_method for p in result.pages] == ["direct", "ocr"]
    assert len(result.tables) == 2
    page0_values = _cells_by_pos(result.tables[0])
    assert page0_values[(1, 0)] == "Товар А"
    page1_values = _cells_by_pos(result.tables[1])
    assert any("Товар Б" in v for v in page1_values.values())


def test_scenario_15_corrupt_pdf_raises_russian_error(tmp_path):
    pdf_path = scenario_15_corrupt_pdf(tmp_path)
    with pytest.raises(CorruptPdfError) as exc_info:
        convert_document(str(pdf_path), ConversionSettings())
    assert "повреждён" in exc_info.value.user_message


def test_scenario_16_encrypted_pdf_raises_russian_error(tmp_path):
    pdf_path = scenario_16_encrypted_pdf(tmp_path)
    with pytest.raises(EncryptedPdfError) as exc_info:
        convert_document(str(pdf_path), ConversionSettings())
    assert "паролем" in exc_info.value.user_message
