from __future__ import annotations

from openpyxl import load_workbook

from pdfxlsx.core.config import ConversionSettings
from pdfxlsx.core.models import DocumentResult, PageOutcome, TypedCell, TypedTable
from pdfxlsx.core.xlsx_writer import write_document


def _cell(value, row, col, row_span=1, col_span=1, is_uncertain=False, comment=None, is_header=False, number_format=None):
    return TypedCell(
        value=value,
        display_text=str(value),
        number_format=number_format,
        row=row,
        col=col,
        row_span=row_span,
        col_span=col_span,
        is_uncertain=is_uncertain,
        comment=comment,
        is_header=is_header,
    )


def _simple_table(page_index=0, table_index=0, source="lines"):
    cells = [
        _cell("A", 0, 0, is_header=True),
        _cell("B", 0, 1, is_header=True),
        _cell(1, 1, 0, number_format="0"),
        _cell(2, 1, 1, number_format="0"),
    ]
    return TypedTable(n_rows=2, n_cols=2, cells=cells, page_index=page_index, table_index_on_page=table_index, source=source)


def test_per_page_mode_creates_one_sheet_per_page_with_table(tmp_path):
    table = _simple_table(page_index=0)
    result = DocumentResult(
        source_pdf="x.pdf",
        page_count=1,
        pages=[PageOutcome(page_index=0, page_kind="text", extraction_method="direct", tables=[table])],
        tables=[table],
    )
    settings = ConversionSettings(sheet_per_page=True, sheet_per_table=False)
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    wb = load_workbook(out)
    assert wb.sheetnames == ["Стр. 1"]
    ws = wb["Стр. 1"]
    assert ws.cell(row=1, column=1).value == "A"
    assert ws.cell(row=1, column=2).value == "B"
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == 2


def test_per_table_mode_creates_one_sheet_per_table(tmp_path):
    t1 = _simple_table(page_index=0, table_index=0)
    t2 = _simple_table(page_index=0, table_index=1)
    result = DocumentResult(
        source_pdf="x.pdf",
        page_count=1,
        pages=[PageOutcome(page_index=0, page_kind="text", extraction_method="direct", tables=[t1, t2])],
        tables=[t1, t2],
    )
    settings = ConversionSettings(sheet_per_page=False, sheet_per_table=True)
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    wb = load_workbook(out)
    assert wb.sheetnames == ["Стр.1_Табл.1", "Стр.1_Табл.2"]


def test_merged_cell_written_correctly(tmp_path):
    cells = [_cell("Header", 0, 0, col_span=2, is_header=True), _cell("a", 1, 0), _cell("b", 1, 1)]
    table = TypedTable(n_rows=2, n_cols=2, cells=cells, page_index=0, table_index_on_page=0, source="lines")
    result = DocumentResult(
        source_pdf="x.pdf",
        page_count=1,
        pages=[PageOutcome(page_index=0, page_kind="text", extraction_method="direct", tables=[table])],
        tables=[table],
    )
    settings = ConversionSettings()
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    wb = load_workbook(out)
    ws = wb.active
    assert "A1:B1" in [str(r) for r in ws.merged_cells.ranges]


def test_uncertain_cell_gets_yellow_fill_and_comment(tmp_path):
    cells = [_cell("42", 0, 0, is_uncertain=True, comment="Требуется ручная проверка (уверенность распознавания: 50%)")]
    table = TypedTable(n_rows=1, n_cols=1, cells=cells, page_index=0, table_index_on_page=0, source="ocr-lines")
    result = DocumentResult(
        source_pdf="x.pdf",
        page_count=1,
        pages=[PageOutcome(page_index=0, page_kind="scanned", extraction_method="ocr", tables=[table])],
        tables=[table],
    )
    settings = ConversionSettings()
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    wb = load_workbook(out)
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    assert cell.fill.start_color.rgb in ("FFFFFF00", "00FFFF00")
    assert cell.comment is not None
    assert "Требуется ручная проверка" in cell.comment.text


def test_uncertain_cell_comment_vml_uses_excel_compatible_prefixes(tmp_path):
    """openpyxl writes VML comment shapes with auto-generated ns0:/ns1:/ns2:
    namespace prefixes, which Excel's legacy VML reader does not recognise
    (it matches prefixes literally rather than by namespace URI), making
    the file unreadable. write_document must rewrite them to the literal
    v:/o:/x: prefixes real Excel/Microsoft-authored VML always uses.
    """
    import zipfile

    cells = [_cell("42", 0, 0, is_uncertain=True, comment="Требуется ручная проверка (уверенность распознавания: 50%)")]
    table = TypedTable(n_rows=1, n_cols=1, cells=cells, page_index=0, table_index_on_page=0, source="ocr-lines")
    result = DocumentResult(
        source_pdf="x.pdf",
        page_count=1,
        pages=[PageOutcome(page_index=0, page_kind="scanned", extraction_method="ocr", tables=[table])],
        tables=[table],
    )
    settings = ConversionSettings()
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    with zipfile.ZipFile(out) as archive:
        vml = archive.read("xl/drawings/commentsDrawing1.vml").decode("utf-8")

    assert "ns0:" not in vml and "ns1:" not in vml and "ns2:" not in vml
    assert 'xmlns:v="urn:schemas-microsoft-com:vml"' in vml
    assert 'xmlns:o="urn:schemas-microsoft-com:office:office"' in vml
    assert 'xmlns:x="urn:schemas-microsoft-com:office:excel"' in vml
    assert "<v:shape" in vml
    assert "<x:ClientData" in vml


def test_vml_fix_does_not_downgrade_file_permissions(tmp_path):
    """The VML-prefix fix rewrites the saved workbook by building a temp
    file via tempfile.mkstemp() and os.replace()-ing it over the output.
    mkstemp() defaults to mode 0600 (owner-only), which os.replace() would
    otherwise carry over onto the final .xlsx, leaving it unreadable by
    anything other than the process that wrote it.
    """
    import stat

    cells = [_cell("42", 0, 0, is_uncertain=True, comment="Требуется ручная проверка")]
    table = TypedTable(n_rows=1, n_cols=1, cells=cells, page_index=0, table_index_on_page=0, source="ocr-lines")
    result = DocumentResult(
        source_pdf="x.pdf",
        page_count=1,
        pages=[PageOutcome(page_index=0, page_kind="scanned", extraction_method="ocr", tables=[table])],
        tables=[table],
    )
    settings = ConversionSettings()
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    plain_table = _simple_table(page_index=0)
    plain_result = DocumentResult(
        source_pdf="y.pdf",
        page_count=1,
        pages=[PageOutcome(page_index=0, page_kind="text", extraction_method="direct", tables=[plain_table])],
        tables=[plain_table],
    )
    plain_out = tmp_path / "plain.xlsx"
    write_document(plain_result, settings, str(plain_out))

    assert stat.S_IMODE(out.stat().st_mode) == stat.S_IMODE(plain_out.stat().st_mode)


def test_highlight_disabled_suppresses_fill(tmp_path):
    cells = [_cell("42", 0, 0, is_uncertain=True, comment="Требуется ручная проверка")]
    table = TypedTable(n_rows=1, n_cols=1, cells=cells, page_index=0, table_index_on_page=0, source="ocr-lines")
    result = DocumentResult(
        source_pdf="x.pdf",
        page_count=1,
        pages=[PageOutcome(page_index=0, page_kind="scanned", extraction_method="ocr", tables=[table])],
        tables=[table],
    )
    settings = ConversionSettings(highlight_uncertain_cells=False)
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    wb = load_workbook(out)
    ws = wb.active
    cell = ws.cell(row=1, column=1)
    assert cell.fill.fill_type is None


def test_grid_source_gets_borders_borderless_does_not(tmp_path):
    grid_table = _simple_table(page_index=0, source="lines")
    text_table = _simple_table(page_index=1, source="text")
    result = DocumentResult(
        source_pdf="x.pdf",
        page_count=2,
        pages=[
            PageOutcome(page_index=0, page_kind="text", extraction_method="direct", tables=[grid_table]),
            PageOutcome(page_index=1, page_kind="text", extraction_method="direct", tables=[text_table]),
        ],
        tables=[grid_table, text_table],
    )
    settings = ConversionSettings()
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    wb = load_workbook(out)
    grid_ws = wb["Стр. 1"]
    text_ws = wb["Стр. 2"]
    assert grid_ws.cell(row=1, column=1).border.left.style == "thin"
    assert text_ws.cell(row=1, column=1).border.left.style is None


def test_leading_zero_text_value_preserved(tmp_path):
    cells = [_cell("007", 0, 0)]
    table = TypedTable(n_rows=1, n_cols=1, cells=cells, page_index=0, table_index_on_page=0, source="text")
    result = DocumentResult(
        source_pdf="x.pdf",
        page_count=1,
        pages=[PageOutcome(page_index=0, page_kind="text", extraction_method="direct", tables=[table])],
        tables=[table],
    )
    settings = ConversionSettings()
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    wb = load_workbook(out)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "007"


def test_plain_text_paragraphs_written_below_table(tmp_path):
    table = _simple_table(page_index=0)
    page = PageOutcome(
        page_index=0,
        page_kind="text",
        extraction_method="direct",
        tables=[table],
        plain_text_paragraphs=["Примечание: см. приложение"],
    )
    result = DocumentResult(source_pdf="x.pdf", page_count=1, pages=[page], tables=[table])
    settings = ConversionSettings()
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    wb = load_workbook(out)
    ws = wb.active
    assert ws.cell(row=5, column=1).value == "Примечание: см. приложение"


def test_no_content_produces_placeholder_sheet(tmp_path):
    result = DocumentResult(source_pdf="x.pdf", page_count=1, pages=[PageOutcome(page_index=0, page_kind="blank", extraction_method="blank")])
    settings = ConversionSettings()
    out = tmp_path / "out.xlsx"
    write_document(result, settings, str(out))

    wb = load_workbook(out)
    assert len(wb.sheetnames) == 1
    assert wb.active.cell(row=1, column=1).value
